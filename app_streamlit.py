import streamlit as st
import pdfplumber
import openpyxl
import re
import pandas as pd
from io import BytesIO

st.set_page_config(page_title="CAD PDF vs Excel Marker Checker", layout="wide")
st.title("🔍 Multi-PDF vs Excel Marker Checker & Verification")

# Inisialisasi Session State untuk menyimpan seluruh hasil pengecekan antar komponen
if "all_check_results" not in st.session_state:
    st.session_state["all_check_results"] = {}

# --- FUNGSI HELPER UNTUK MEMBUAT FILE EXCEL DOWNLOAD ---
def convert_df_to_excel(df):
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Hasil_Checking')
    return output.getvalue()

def convert_all_to_excel(all_results):
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Tab 1: Akumulasi Semua Result
        all_rows = []
        for k, rows in all_results.items():
            all_rows.extend(rows)
        if all_rows:
            df_all = pd.DataFrame(all_rows)
            df_all.to_excel(writer, index=False, sheet_name='Ringkasan_Semua_Panel')
        
        # Tab Berikuntnya: Per Komponen
        for key_name, rows in all_results.items():
            if rows:
                df_comp = pd.DataFrame(rows)
                # Pembersihan nama sheet (maksimal 31 karakter & tanpa karakter terlarang)
                sheet_title = key_name[:30].replace("[", "").replace("]", "").replace("*", "").replace(":", "").replace("?", "").replace("/", "")
                df_comp.to_excel(writer, index=False, sheet_name=sheet_title)
                
    return output.getvalue()


# --- 1. PARSER PDF MARKER DENGAN CACHING ---
@st.cache_data(show_spinner=False)
def parse_pdf_bytes(file_bytes, filename):
    extracted_data = {
        "filename": filename,
        "ratio": None,
        "length_yds": None,   
        "length_inch": None,  
        "width": None        
    }
    
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        text = ""
        for page in pdf.pages:
            text += (page.extract_text() or "") + "\n"
            
    # Extract Ratio
    ratio_match = re.search(r"Sizes\s+Ratio\s*:\s*(.+)", text, re.IGNORECASE)
    if ratio_match:
        raw_ratio = ratio_match.group(1).strip()
        extracted_data["ratio"] = raw_ratio.replace("*", "/").replace(" ", "")

    # Extract Width
    width_match = re.search(r"Marker\s+Width\s*:\s*([0-9\.\"]+)", text, re.IGNORECASE)
    if width_match:
        val = width_match.group(1).replace('"', '').strip()
        extracted_data["width"] = float(val) if val else None

    # Extract Length
    length_yd_match = re.search(r"Marker\s+Length\s*:\s*(\d+)\s*yd\s*([\d\.]+)\s*\"", text, re.IGNORECASE)
    if length_yd_match:
        yds = int(length_yd_match.group(1))
        inch = float(length_yd_match.group(2))
        extracted_data["length_yds"] = yds
        extracted_data["length_inch"] = inch
    else:
        length_inch_match = re.search(r"Marker\s+Length\s*:\s*([\d\.]+)\s*\"", text, re.IGNORECASE)
        if length_inch_match:
            extracted_data["length_yds"] = 0 
            extracted_data["length_inch"] = float(length_inch_match.group(1))
        
    return extracted_data


# --- 2. UPLOAD FILE EXCEL ---
excel_file = st.file_uploader("1. Unggah Master Excel yang Sudah Terisi (.xlsx)", type=["xlsx"], key="excel_uploader")

if excel_file:
    wb = openpyxl.load_workbook(BytesIO(excel_file.getvalue()), data_only=True)
    selected_sheet_name = st.selectbox("Pilih Tab/Sheet Excel:", options=wb.sheetnames)
    ws = wb[selected_sheet_name]

    # --- 3. DETEKSI KOMPONEN & DATA EXCEL ---
    all_components = []
    current_comp = None

    for r in range(1, 400):
        row_cells = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, 11)]
        row_text_combined = " ".join(row_cells).lower()

        if "component" in row_text_combined:
            comp_title = None
            for col_idx in range(2, 9):
                val = ws.cell(row=r, column=col_idx).value
                if val and str(val).strip() != "" and str(val).strip().lower() != "component":
                    comp_title = str(val).strip()
                    break
            
            if comp_title:
                current_comp = {
                    "component_name": comp_title.replace("\n", " "),
                    "start_row": r,
                    "ratios": []
                }
                all_components.append(current_comp)
                continue

        if current_comp:
            if any("marker" in cell.lower() or "ratio" in cell.lower() for cell in row_cells):
                ratio_val = None
                for c_idx in range(2, 6):
                    v = str(ws.cell(row=r, column=c_idx).value or "").strip()
                    if v and not any(k in v.lower() for k in ["color", "code", "fabric", "garment"]):
                        ratio_val = v
                        break
                
                if ratio_val:
                    current_comp["ratios"].append({
                        "ratio_text": ratio_val.replace(" ", ""),
                        "row": r  
                    })

    st.info(f"Terdeteksi **{len(all_components)}** Komponen pada Tab **{selected_sheet_name}**")

    # --- 4. PILIHAN KOMPONEN TARGET ---
    if all_components:
        comp_options = [c["component_name"] for c in all_components]
        selected_comp_name = st.selectbox("🎯 Pilih Komponen Target yang Ingin Dicek:", options=comp_options)
        selected_comp = next(c for c in all_components if c["component_name"] == selected_comp_name)

        # --- 5. KONFIGURASI KOLOM EXCEL & ALLOWANCE INCH ---
        col_k1, col_k2, col_k3, col_k4 = st.columns(4)
        col_yds = col_k1.number_input("Kolom Length (Yds)", min_value=1, value=15)
        col_inch = col_k2.number_input("Kolom Length (Inch)", min_value=1, value=16)
        col_width = col_k3.number_input("Kolom Width", min_value=1, value=17)
        inch_allowance = col_k4.number_input("Tambahan Inch (Allowance)", value=2.0, step=0.5, help="Nilai inch PDF akan ditambahkan angka ini sebelum dicocokkan ke Excel")

        # --- 6. UPLOAD PDF DENGAN KEY DINAMIS PER KOMPONEN ---
        pdf_files = st.file_uploader(
            f"2. Unggah File CAD PDF Marker khusus untuk [{selected_comp_name}]", 
            type=["pdf"], 
            accept_multiple_files=True,
            key=f"pdf_uploader_{selected_sheet_name}_{selected_comp_name}"
        )

        # --- 7. PROSES PENGECEKAN / COMPARISON ---
        if pdf_files:
            comparison_results = []
            unmatched_pdfs = []

            for pdf_file in pdf_files:
                pdf_data = parse_pdf_bytes(pdf_file.getvalue(), pdf_file.name)
                raw_pdf_ratio = (pdf_data["ratio"] or "").upper().replace(" ", "").replace("*", "/")
                
                found_in_excel = False
                for item in selected_comp["ratios"]:
                    raw_excel_ratio = item["ratio_text"].upper().replace(" ", "").replace("*", "/")
                    
                    if raw_pdf_ratio and (raw_pdf_ratio == raw_excel_ratio):
                        found_in_excel = True
                        t_row = item["row"]

                        # Ambil nilai Excel
                        excel_yds = ws.cell(row=t_row, column=int(col_yds)).value
                        excel_inch = ws.cell(row=t_row, column=int(col_inch)).value
                        excel_width = ws.cell(row=t_row, column=int(col_width)).value

                        # Clean nilai Excel
                        try:
                            excel_yds_num = int(excel_yds) if excel_yds not in [None, ""] else 0
                        except:
                            excel_yds_num = 0

                        try:
                            excel_inch_num = round(float(excel_inch), 2) if excel_inch not in [None, ""] else 0.0
                        except:
                            excel_inch_num = 0.0

                        try:
                            excel_width_num = round(float(excel_width), 2) if excel_width not in [None, ""] else 0.0
                        except:
                            excel_width_num = 0.0

                        # Nilai PDF
                        pdf_yds_num = int(pdf_data["length_yds"] or 0)
                        raw_pdf_inch_num = round(float(pdf_data["length_inch"] or 0.0), 2)
                        pdf_width_num = round(float(pdf_data["width"] or 0.0), 2)

                        # Perhitungan PDF + Allowance (+ 2 inch)
                        expected_pdf_inch = round(raw_pdf_inch_num + float(inch_allowance), 2)

                        # Evaluasi
                        match_ratio = True
                        match_yds = (excel_yds_num == pdf_yds_num)
                        match_inch = abs(excel_inch_num - expected_pdf_inch) < 0.01
                        match_width = abs(excel_width_num - pdf_width_num) < 0.01

                        all_matched = match_ratio and match_yds and match_inch and match_width

                        if all_matched:
                            status = "🟢 MATCH"
                            notes = "Semua Sesuai (Ratio, Yds, Inch+2\", Width)"
                        else:
                            status = "🔴 MISMATCH"
                            diffs = []
                            if not match_yds:
                                diffs.append(f"Yds Beda (Excel: {excel_yds_num} vs PDF: {pdf_yds_num})")
                            if not match_inch:
                                diffs.append(f"Inch Beda (Excel: {excel_inch_num} vs Expected PDF+Allow: {expected_pdf_inch})")
                            if not match_width:
                                diffs.append(f"Width Beda (Excel: {excel_width_num} vs PDF: {pdf_width_num})")
                            notes = " | ".join(diffs)

                        comparison_results.append({
                            "Komponen": selected_comp_name,
                            "Status Overall": status,
                            "Nama PDF": pdf_data["filename"],
                            "Baris Excel": t_row,
                            "Ratio (Excel / PDF)": f"{item['ratio_text']} / {pdf_data['ratio']} 🟢",
                            "Yds (Excel / PDF)": f"{excel_yds_num} / {pdf_yds_num} " + ("🟢" if match_yds else "🔴"),
                            "Inch (Excel / Expected PDF+2\")": f"{excel_inch_num} / {expected_pdf_inch} (Asli: {raw_pdf_inch_num}) " + ("🟢" if match_inch else "🔴"),
                            "Width (Excel / PDF)": f"{excel_width_num} / {pdf_width_num} " + ("🟢" if match_width else "🔴"),
                            "Keterangan": notes
                        })
                        break
                
                if not found_in_excel:
                    unmatched_pdfs.append({
                        "Komponen": selected_comp_name,
                        "Nama PDF": pdf_data["filename"],
                        "Ratio PDF": pdf_data["ratio"] or "Tidak Terbaca",
                        "Catatan": "Rasio PDF tidak ditemukan di daftar ratio Excel komponen ini"
                    })

            # SIMPAN HASIL KOMPONEN INI KE SESSION STATE
            st.session_state["all_check_results"][f"{selected_sheet_name}_{selected_comp_name}"] = comparison_results

        # --- 8. TAMPILKAN HASIL KOMPONEN SAAT INI ---
        current_key = f"{selected_sheet_name}_{selected_comp_name}"
        if current_key in st.session_state["all_check_results"] and st.session_state["all_check_results"][current_key]:
            st.markdown("---")
            col_h1, col_h2 = st.columns([3, 1])
            with col_h1:
                st.subheader(f"📊 Hasil Pengecekan Komponen: [{selected_comp_name}]")
            
            df_curr = pd.DataFrame(st.session_state["all_check_results"][current_key])
            total_check = len(df_curr)
            total_match = len(df_curr[df_curr["Status Overall"] == "🟢 MATCH"])
            total_mismatch = len(df_curr[df_curr["Status Overall"] == "🔴 MISMATCH"])

            # TOMBOL DOWNLOAD UNTUK KOMPONEN INI
            excel_curr_bytes = convert_df_to_excel(df_curr)
            with col_h2:
                st.download_button(
                    label=f"📥 Download Excel ({selected_comp_name})",
                    data=excel_curr_bytes,
                    file_name=f"Hasil_Check_{selected_comp_name}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

            m1, m2, m3 = st.columns(3)
            m1.metric("PDF Dicek (Komponen Ini)", total_check)
            m2.metric("Sesuai (Match)", total_match)
            m3.metric("Beda (Mismatch)", total_mismatch, delta_color="inverse")

            st.dataframe(df_curr, use_container_width=True)

        # --- 9. TAMPILKAN REKAPITULASI SEMUA KOMPONEN / PANEL ---
        if st.session_state["all_check_results"]:
            all_rows = []
            for k, rows in st.session_state["all_check_results"].items():
                all_rows.extend(rows)
            
            if all_rows:
                st.markdown("---")
                col_r1, col_r2 = st.columns([3, 1])
                with col_r1:
                    st.subheader("📋 Ringkasan Akumulasi Pengecekan (Semua Panel/Komponen)")
                
                df_all = pd.DataFrame(all_rows)

                # TOMBOL DOWNLOAD UNTUK SEMUA REKAP
                excel_all_bytes = convert_all_to_excel(st.session_state["all_check_results"])
                with col_r2:
                    st.download_button(
                        label="📥 Download Excel ALL Panel (Multi-Tab)",
                        data=excel_all_bytes,
                        file_name="Rekap_Hasil_Check_Semua_Panel.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True
                    )

                tot_all = len(df_all)
                tot_match_all = len(df_all[df_all["Status Overall"] == "🟢 MATCH"])
                tot_mismatch_all = len(df_all[df_all["Status Overall"] == "🔴 MISMATCH"])

                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Total Komponen Dicek", len(st.session_state["all_check_results"]))
                c2.metric("Total PDF Akumulasi", tot_all)
                c3.metric("Total Sesuai (🟢 MATCH)", tot_match_all)
                c4.metric("Total Beda (🔴 MISMATCH)", tot_mismatch_all, delta_color="inverse")

                st.dataframe(df_all, use_container_width=True)

                if st.button("🗑️ Reset / Hapus Semua Memory Pengecekan"):
                    st.session_state["all_check_results"] = {}
                    st.rerun()